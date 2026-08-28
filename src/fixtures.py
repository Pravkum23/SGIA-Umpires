import csv
import re
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel
from sqlalchemy import text

FIXTURE_COLUMNS = ["Day", "Date", "Time", "TEAM 1", "TEAM 2"]
CURRENT_FIXTURE_COLUMNS = FIXTURE_COLUMNS + ["Poll State", "Match Status"]
PREVIEW_COLUMNS = ["Line", *FIXTURE_COLUMNS, "Status", "Error"]
TEMPLATE_ROWS = [
    ["Saturday", "15-Aug-2026", "11:00 AM", "Changi Risers", "Black Panthers"],
    ["Saturday", "15-Aug-2026", "3:00 PM", "Knights United", "Redbacks"],
]

HEADER_ALIASES = {
    "day": "Day",
    "date": "Date",
    "time": "Time",
    "team1": "TEAM 1",
    "hometeam": "TEAM 1",
    "team2": "TEAM 2",
    "awayteam": "TEAM 2",
}
REQUIRED_COLUMNS = {"Date", "Time", "TEAM 1", "TEAM 2"}


def _header_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _excel_datetime_value(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value)
        except (TypeError, ValueError, OverflowError):
            return value
    return value


def _parse_date(value):
    value = _excel_datetime_value(value)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _clean_text(value)
    for pattern in ("%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    for pattern in ("%d-%b-%y", "%d/%m/%y", "%d-%m-%y", "%y-%m-%d"):
        try:
            parsed = datetime.strptime(raw, pattern).date()
            # Spreadsheet schedules use two-digit years for the current
            # century (for example 01-Sep-26 means 1 September 2026).
            return parsed.replace(year=parsed.year + 100) if parsed.year < 2000 else parsed
        except ValueError:
            continue
    raise ValueError(f"Unsupported date: {raw or 'blank'}")


def _parse_time(value):
    value = _excel_datetime_value(value)
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    if isinstance(value, timedelta):
        seconds = int(value.total_seconds()) % (24 * 3600)
        return time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)
    raw = _clean_text(value).upper()
    for pattern in ("%I:%M %p", "%I:%M:%S %p", "%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(raw, pattern).time()
        except ValueError:
            continue
    raise ValueError(f"Unsupported time: {raw or 'blank'}")


def _parse_datetime(date_value, time_value):
    # Fixture values are Singapore wall-clock values. Keep the result naive so
    # neither the server nor the browser can shift it to another timezone.
    return datetime.combine(_parse_date(date_value), _parse_time(time_value))


def _display_datetime(starts_at):
    if not isinstance(starts_at, datetime):
        starts_at = datetime.fromisoformat(str(starts_at))
    return {
        "Day": starts_at.strftime("%A"),
        "Date": starts_at.strftime("%d-%b-%Y"),
        "Time": starts_at.strftime("%I:%M %p").lstrip("0"),
    }


def add_fixture(engine, date_value, time_value, team_1, team_2):
    starts_at = _parse_datetime(date_value, time_value)
    team_1, team_2 = _clean_text(team_1), _clean_text(team_2)
    if not team_1 or not team_2:
        raise ValueError("Both teams are required")
    with engine.begin() as connection:
        result = connection.execute(text("""
            INSERT INTO fixtures(starts_at,home_team,away_team,availability_open,poll_state)
            VALUES(:starts,:home,:away,false,'CLOSED') ON CONFLICT(starts_at) DO NOTHING
        """), {"starts": starts_at, "home": team_1, "away": team_2})
        return result.rowcount == 1


def edit_fixture(engine, fixture_id, date_value, time_value, team_1, team_2):
    starts_at = _parse_datetime(date_value, time_value)
    team_1, team_2 = _clean_text(team_1), _clean_text(team_2)
    if not team_1 or not team_2:
        raise ValueError("Both teams are required")
    with engine.begin() as connection:
        result = connection.execute(text("""
            UPDATE fixtures SET starts_at=:starts,home_team=:home,away_team=:away WHERE id=:fixture
        """), {"starts": starts_at, "home": team_1, "away": team_2, "fixture": fixture_id})
        return result.rowcount == 1


def _delimiter(text_value):
    first_line = next((line for line in text_value.splitlines() if line.strip()), "")
    if "|" in first_line:
        return "|"
    if "\t" in first_line:
        return "\t"
    if "," in first_line:
        return ","
    if ";" in first_line:
        return ";"
    return None


def _raw_rows_from_values(values):
    populated = [(index, list(row)) for index, row in enumerate(values, 1) if any(_clean_text(cell) for cell in row)]
    if not populated:
        return [{"Line": 1, "_source_error": "No fixture rows found"}]

    first_number, first = populated[0]
    mapped_headers = [HEADER_ALIASES.get(_header_key(value)) for value in first]
    has_header = sum(header is not None for header in mapped_headers) >= 2
    if has_header:
        present = {header for header in mapped_headers if header}
        missing = REQUIRED_COLUMNS - present
        if missing:
            names = ", ".join(sorted(missing))
            return [{"Line": first_number, "_source_error": f"Missing required column(s): {names}"}]
        source_rows = populated[1:]
        result = []
        for line_number, row in source_rows:
            item = {"Line": line_number}
            for index, header in enumerate(mapped_headers):
                if header:
                    item[header] = row[index] if index < len(row) else None
            result.append(item)
        return result or [{"Line": first_number + 1, "_source_error": "No fixture rows found below the header"}]

    result = []
    for line_number, row in populated:
        if len(row) == 5:
            item = dict(zip(FIXTURE_COLUMNS, row))
        elif len(row) == 4:
            item = dict(zip(FIXTURE_COLUMNS[1:], row))
        else:
            item = {"_source_error": "Expected 4 columns (Date, Time, Team 1, Team 2) or 5 columns including Day"}
        item["Line"] = line_number
        result.append(item)
    return result


def _raw_rows_from_text(text_value):
    delimiter = _delimiter(text_value)
    if not delimiter:
        return [{"Line": 1, "_source_error": "Use pipe, tab, or CSV-separated columns; spaces cannot safely separate team names"}]
    return _raw_rows_from_values(csv.reader(StringIO(text_value), delimiter=delimiter, skipinitialspace=True))


def _preview_rows(engine, raw_rows):
    with engine.connect() as connection:
        existing = {
            row.starts_at if isinstance(row.starts_at, datetime) else datetime.fromisoformat(str(row.starts_at))
            for row in connection.execute(text("SELECT starts_at FROM fixtures"))
        }
    seen = set(existing)
    preview = []
    for raw in raw_rows:
        item = {column: "" for column in PREVIEW_COLUMNS}
        item["Line"] = raw.get("Line", len(preview) + 1)
        item["Status"] = "ERROR"
        item.update({column: _clean_text(raw.get(column)) for column in FIXTURE_COLUMNS})
        try:
            if raw.get("_source_error"):
                raise ValueError(raw["_source_error"])
            starts_at = _parse_datetime(raw.get("Date"), raw.get("Time"))
            team_1, team_2 = _clean_text(raw.get("TEAM 1")), _clean_text(raw.get("TEAM 2"))
            if not team_1 or not team_2:
                raise ValueError("Both teams are required")
            item.update(_display_datetime(starts_at))
            item.update({"TEAM 1": team_1, "TEAM 2": team_2, "starts_at": starts_at})
            if starts_at in seen:
                item["Status"] = "DUPLICATE"
                item["Error"] = "A fixture already uses this Singapore date/time"
            else:
                item["Status"] = "READY"
            seen.add(starts_at)
        except (TypeError, ValueError) as error:
            item["Error"] = str(error)
        preview.append(item)
    return preview


def preview_bulk_fixtures(engine, pasted_text):
    """Preview pipe, tab, or CSV text using the shared safe import path."""
    return _preview_rows(engine, _raw_rows_from_text(pasted_text or ""))


def preview_fixture_file(engine, filename, content):
    suffix = str(filename or "").lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        text_value = content.decode("utf-8-sig") if isinstance(content, bytes) else str(content)
        raw_rows = _raw_rows_from_text(text_value)
    elif suffix == "xlsx":
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        raw_rows = _raw_rows_from_values(workbook.active.iter_rows(values_only=True))
    else:
        raw_rows = [{"Line": 1, "_source_error": "Only .csv and .xlsx files are supported"}]
    return _preview_rows(engine, raw_rows)


def import_bulk_fixtures(engine, preview):
    imported = 0
    with engine.begin() as connection:
        for item in preview:
            if item.get("Status") != "READY":
                continue
            result = connection.execute(text("""
                INSERT INTO fixtures(starts_at,home_team,away_team,availability_open,poll_state)
                VALUES(:starts,:home,:away,false,'CLOSED') ON CONFLICT(starts_at) DO NOTHING
            """), {"starts": item["starts_at"], "home": item["TEAM 1"], "away": item["TEAM 2"]})
            imported += max(result.rowcount, 0)
    return imported


def fixture_template_csv():
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(FIXTURE_COLUMNS)
    writer.writerows(TEMPLATE_ROWS)
    return output.getvalue().encode("utf-8-sig")


def _styled_workbook(headers, records, example_rows=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fixtures"
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F78")
    for record in records:
        sheet.append([record.get(header, "") for header in headers])
    if example_rows:
        fill = PatternFill("solid", fgColor="FFF2CC")
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.fill = fill
                cell.font = Font(italic=True)
            row[0].comment = Comment("Example row — replace or delete before uploading your schedule.", "SGIA Umpires")
    widths = {"Day": 14, "Date": 16, "Time": 13, "TEAM 1": 24, "TEAM 2": 24, "Poll State": 15, "Match Status": 16}
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[chr(64 + index)].width = widths.get(header, 18)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def fixture_template_xlsx():
    records = []
    for day_value, date_value, time_value, team_1, team_2 in TEMPLATE_ROWS:
        records.append({
            "Day": day_value,
            "Date": _parse_date(date_value),
            "Time": _parse_time(time_value),
            "TEAM 1": team_1,
            "TEAM 2": team_2,
        })
    content = _styled_workbook(FIXTURE_COLUMNS, records, example_rows=True)
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    for cell in sheet["B"][1:]:
        cell.number_format = "DD-MMM-YYYY"
    for cell in sheet["C"][1:]:
        cell.number_format = "h:mm AM/PM"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def current_fixture_records(engine):
    with engine.connect() as connection:
        fixtures = list(connection.execute(text("""
            SELECT starts_at,home_team,away_team,poll_state,match_status
            FROM fixtures ORDER BY starts_at
        """)).mappings())
    records = []
    for fixture in fixtures:
        display = _display_datetime(fixture.starts_at)
        records.append({
            **display,
            "TEAM 1": fixture.home_team,
            "TEAM 2": fixture.away_team,
            "Poll State": fixture.poll_state,
            "Match Status": fixture.match_status,
        })
    return records


def current_fixtures_csv(engine):
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=CURRENT_FIXTURE_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(current_fixture_records(engine))
    return output.getvalue().encode("utf-8-sig")


def current_fixtures_xlsx(engine):
    return _styled_workbook(CURRENT_FIXTURE_COLUMNS, current_fixture_records(engine))
