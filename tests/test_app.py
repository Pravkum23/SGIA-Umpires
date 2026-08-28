from io import BytesIO
from pathlib import Path
from datetime import date, datetime, time

from openpyxl import Workbook, load_workbook
from PIL import Image
from sqlalchemy import create_engine, text
from streamlit.testing.v1 import AppTest

from src.allocation import allocation_message, propose
from src.allocation_scope import current_allocation_summary, published_allocation_cycles
from src.board import BOARD_COLUMNS, REVIEW_COLUMNS, _font, allocation_board, allocation_board_csv, allocation_board_png, allocation_review, confirm_all_proposed, confirmed_allocation_board, save_allocation_board, save_allocation_review
from src.db import POSTGRES_SCHEMA, SQLITE_SCHEMA, initialize, people_seed_params, rows, save_vote, schema_for
from src.fixtures import (
    CURRENT_FIXTURE_COLUMNS,
    FIXTURE_COLUMNS,
    current_fixtures_csv,
    current_fixtures_xlsx,
    add_fixture,
    edit_fixture,
    fixture_template_csv,
    fixture_template_xlsx,
    import_bulk_fixtures,
    preview_bulk_fixtures,
    preview_fixture_file,
)
from src.lifecycle import authenticate_person, availability_for_person, forget_person, person_for_token, remember_person, replace_assignment, set_person_pin, suggest_replacement, verify_person_pin, withdraw_assignment
from src.live_poll import coverage_status, live_poll_monitor
from src.poll_audit import activity_history, coarse_client, record_poll_view, response_audit, response_audit_csv, set_fixture_poll_state, transition_poll_state, update_open_slots
from src.routing import is_admin_request
from src.workload import fixture_history, season_workload, set_match_status


def db():
    engine = create_engine("sqlite:///:memory:")
    initialize(engine)
    return engine


def open_and_vote_all(engine, fixture_ids):
    with engine.begin() as connection:
        for fixture_id in fixture_ids:
            connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id=:id"), {"id": fixture_id})
            connection.execute(text("INSERT INTO availability(person_id,fixture_id) SELECT id,:id FROM people"), {"id": fixture_id})


def test_seed_is_idempotent_and_exact():
    engine = db()
    initialize(engine)
    assert rows(engine, "SELECT COUNT(*) n FROM fixtures")[0]["n"] == 29
    assert rows(engine, "SELECT COUNT(*) n FROM people")[0]["n"] == 16


def test_vote_can_change_and_closed_slots_are_not_public():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (1,2)"))
    save_vote(engine, 1, [1, 2])
    save_vote(engine, 1, [2])
    assert rows(engine, "SELECT fixture_id FROM availability WHERE person_id=1") == [{"fixture_id": 2}]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id=2"))
    assert rows(engine, "SELECT id FROM fixtures WHERE availability_open=true") == [{"id": 1}]


def test_new_poll_update_preserves_closed_fixture_votes():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (1,2)"))
    save_vote(engine, 1, [1, 2])
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id IN (1,2)"))
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (12,13)"))
    save_vote(engine, 1, [12])
    assert rows(engine, "SELECT fixture_id FROM availability WHERE person_id=1 ORDER BY fixture_id") == [
        {"fixture_id": 1}, {"fixture_id": 2}, {"fixture_id": 12}
    ]


def test_editing_open_poll_only_adds_and_removes_open_choices():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (1,2)"))
    save_vote(engine, 1, [1, 2])
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id=1"))
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id=3"))
    save_vote(engine, 1, [3])
    assert rows(engine, "SELECT fixture_id FROM availability WHERE person_id=1 ORDER BY fixture_id") == [
        {"fixture_id": 1}, {"fixture_id": 3}
    ]
    events = rows(engine, "SELECT fixture_id,event_type FROM availability_events WHERE person_id=1 ORDER BY id")
    assert events[-2:] == [{"fixture_id": 3, "event_type": "ADDED"}, {"fixture_id": 2, "event_type": "REMOVED"}]


def test_frozen_poll_is_read_only_and_keeps_existing_selection():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id=1"))
    save_vote(engine, 1, [1])
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id=1"))
    save_vote(engine, 1, [])
    assert rows(engine, "SELECT fixture_id FROM availability WHERE person_id=1") == [{"fixture_id": 1}]
    assert rows(engine, "SELECT event_type FROM availability_events WHERE person_id=1") == [{"event_type": "ADDED"}]


def test_latest_active_availability_drives_allocation():
    engine = db()
    malo = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (8,9,10)"))
    save_vote(engine, malo, [8, 9, 10])
    save_vote(engine, malo, [8, 10])
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id IN (8,9,10)"))
    propose(engine)
    assert availability_for_person(engine, malo, "FROZEN") == {8, 10}
    assert not rows(engine, "SELECT 1 FROM assignments WHERE fixture_id=9 AND person_id=:person", {"person": malo})


def test_device_token_recognition_and_change_person():
    engine = db()
    malo = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    token = remember_person(engine, malo)
    assert len(token) >= 32
    assert person_for_token(engine, token) == {"id": malo, "name": "Malo"}
    assert "Malo" not in token
    forget_person(engine, token)
    assert person_for_token(engine, token) is None


def test_wrong_pin_rejected_and_correct_pin_creates_device_token():
    engine = db()
    malo = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    set_person_pin(engine, malo, "4821")
    token, result = authenticate_person(engine, malo, "1111")
    assert token is None and result == "INVALID"
    token, result = authenticate_person(engine, malo, "4821")
    assert result == "OK" and token
    assert person_for_token(engine, token) == {"id": malo, "name": "Malo"}
    stored = rows(engine, "SELECT pin_hash FROM people WHERE id=:person", {"person": malo})[0]["pin_hash"]
    assert "4821" not in stored and stored.startswith("pbkdf2_sha256$")


def test_pin_reset_revokes_existing_device_token():
    engine = db()
    malo = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    set_person_pin(engine, malo, "4821")
    token, _ = authenticate_person(engine, malo, "4821")
    set_person_pin(engine, malo, "7319", revoke_tokens=True)
    assert person_for_token(engine, token) is None
    assert verify_person_pin(engine, malo, "7319") == (True, "OK")


def test_returning_volunteer_ui_is_recognized_and_choices_prepopulate(monkeypatch, tmp_path):
    path = tmp_path / "remembered.db"
    engine = create_engine(f"sqlite:///{path.as_posix()}")
    initialize(engine)
    malo = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=true,poll_state='OPEN' WHERE id IN (1,2)"))
    save_vote(engine, malo, [1])
    token = remember_person(engine, malo)
    engine.dispose()
    monkeypatch.setenv("SGIA_SQLITE_PATH", str(path))
    page = AppTest.from_file(Path(__file__).parents[1] / "app.py")
    page.query_params["device"] = token
    page.run(timeout=20)
    assert not list(page.exception)
    assert any("Hi Malo" in getattr(item, "value", "") for item in page.markdown)
    choices = {item.label: item.value for item in page.checkbox}
    assert choices["Saturday - 11:00 AM"] is True
    assert choices["Saturday - 3:00 PM"] is False


def test_sqlite_and_postgres_schema_use_native_auto_generated_ids():
    assert schema_for("sqlite") == SQLITE_SCHEMA
    assert "INTEGER PRIMARY KEY AUTOINCREMENT" in SQLITE_SCHEMA
    assert schema_for("postgresql") == POSTGRES_SCHEMA
    assert POSTGRES_SCHEMA.count("GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY") == 7
    assert "AUTOINCREMENT" not in POSTGRES_SCHEMA


def test_people_seed_parameters_use_postgresql_safe_booleans():
    params = people_seed_params()
    assert len(params) == 16
    assert all(type(person["u"]) is bool for person in params)
    assert all(type(person["s"]) is bool for person in params)


def test_closed_poll_votes_can_still_be_allocated():
    engine = db()
    open_and_vote_all(engine, [1])
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id=1"))
    propose(engine)
    assert rows(engine, "SELECT COUNT(*) n FROM assignments WHERE fixture_id=1")[0]["n"] == 3


def test_three_role_staffing_model_is_proposed():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    proposed = rows(engine, "SELECT role FROM assignments WHERE fixture_id=1 ORDER BY role")
    assert proposed == [{"role": "scorer"}, {"role": "umpire_1"}, {"role": "umpire_2"}]
    assert rows(engine, "SELECT COUNT(DISTINCT person_id) n FROM assignments WHERE fixture_id=1")[0]["n"] == 3


def test_legacy_umpire_storage_is_migrated_idempotently():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,reason) VALUES(1,'umpire',1,true,'Legacy')"))
    initialize(engine)
    initialize(engine)
    assert rows(engine, "SELECT role,person_id,confirmed FROM assignments") == [
        {"role": "umpire_1", "person_id": 1, "confirmed": 1}
    ]


def test_regenerate_changes_only_unconfirmed_and_preserves_confirmed_exactly():
    engine = db()
    open_and_vote_all(engine, [1, 2])
    propose(engine)
    before = rows(engine, "SELECT fixture_id,role,person_id,confirmed FROM assignments ORDER BY fixture_id,role")
    confirmed = before[0]
    unconfirmed = before[-1]
    with engine.begin() as connection:
        connection.execute(text("UPDATE assignments SET confirmed=true WHERE fixture_id=:fixture_id AND role=:role"), confirmed)
        connection.execute(text("DELETE FROM availability WHERE fixture_id=:fixture AND person_id=:person"), {"fixture": unconfirmed["fixture_id"], "person": unconfirmed["person_id"]})
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='FROZEN' WHERE id IN (1,2)"))
    propose(engine, regenerate_unconfirmed=True)
    kept = rows(engine, "SELECT person_id,confirmed FROM assignments WHERE fixture_id=:fixture_id AND role=:role", confirmed)[0]
    refreshed = rows(engine, "SELECT person_id,confirmed FROM assignments WHERE fixture_id=:fixture_id AND role=:role", unconfirmed)[0]
    assert kept == {"person_id": confirmed["person_id"], "confirmed": 1}
    assert refreshed["person_id"] != unconfirmed["person_id"]
    assert refreshed["confirmed"] == 0


def test_special_role_rules_and_praveen_scorer_preference():
    engine = db()
    open_and_vote_all(engine, [1, 2])
    propose(engine)
    assert not rows(engine, "SELECT 1 FROM assignments a JOIN people p ON p.id=a.person_id WHERE a.role='scorer' AND p.name IN ('Velu','Shree')")
    praveen = rows(engine, "SELECT a.role FROM assignments a JOIN people p ON p.id=a.person_id WHERE p.name='Praveen'")
    assert praveen and all(item["role"] == "scorer" for item in praveen)


def test_weekend_pair_is_back_to_back_opposite_roles_without_third_duty():
    engine = db()
    open_and_vote_all(engine, [1, 2, 3])
    propose(engine)
    paired = rows(engine, """
        SELECT p.name FROM assignments first
        JOIN assignments second ON second.fixture_id=2 AND second.person_id=first.person_id
        JOIN people p ON p.id=first.person_id
        WHERE first.fixture_id=1
          AND ((first.role='scorer' AND second.role IN ('umpire_1','umpire_2'))
            OR (second.role='scorer' AND first.role IN ('umpire_1','umpire_2')))
    """)
    assert paired
    triple = rows(engine, "SELECT person_id,COUNT(*) n FROM assignments WHERE fixture_id IN (1,2,3) GROUP BY person_id HAVING COUNT(*)>2")
    assert not triple


def test_weekday_opposite_role_balance_when_available():
    engine = db()
    open_and_vote_all(engine, [8, 9])  # Tuesday and Wednesday
    propose(engine)
    balanced = rows(engine, """
        SELECT person_id FROM assignments WHERE fixture_id IN (8,9)
        GROUP BY person_id
        HAVING SUM(CASE WHEN role='scorer' THEN 1 ELSE 0 END)>0
           AND SUM(CASE WHEN role IN ('umpire_1','umpire_2') THEN 1 ELSE 0 END)>0
    """)
    assert balanced


def test_allocation_board_columns_and_manual_save():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    board = allocation_board(engine)
    assert list({key: None for key in board[0] if key != "fixture_id"}) == BOARD_COLUMNS
    original = board[0]["Umpire 1"]
    replacement = next(row["name"] for row in rows(engine, "SELECT name FROM people WHERE active=true ORDER BY name") if row["name"] not in {original, board[0]["Umpire 2"], board[0]["Scorer"]})
    board[0]["Umpire 1"] = replacement
    save_allocation_board(engine, board)
    saved = rows(engine, """
        SELECT p.name,a.confirmed,a.reason FROM assignments a JOIN people p ON p.id=a.person_id
        WHERE a.fixture_id=1 AND a.role='umpire_1'
    """)[0]
    assert saved == {"name": replacement, "confirmed": 1, "reason": "Admin allocation board"}
    audit = rows(engine, "SELECT event_type,person_id,replacement_person_id FROM assignment_events WHERE fixture_id=1 AND role='umpire_1'")
    assert audit[-1] == {"event_type": "MANUAL_CHANGE", "person_id": next(row["id"] for row in rows(engine, "SELECT id,name FROM people") if row["name"] == original), "replacement_person_id": next(row["id"] for row in rows(engine, "SELECT id,name FROM people") if row["name"] == replacement)}
    assert rows(engine, "SELECT COUNT(*) n FROM assignment_events WHERE event_type='CONFIRMED'")[0]["n"] == 2


def test_allocation_board_png_export():
    engine = db()
    open_and_vote_all(engine, [1, 2])
    propose(engine)
    image = allocation_board_png(allocation_board(engine), Path(__file__).parents[1] / "assets" / "sgia-logo.png")
    assert image.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(image) > 10_000
    title_bounds = _font(52, True).getbbox("SGIA Official")
    cell_bounds = _font(24, True).getbbox("Black Panthers")
    assert title_bounds[3] - title_bounds[1] >= 38
    assert cell_bounds[3] - cell_bounds[1] >= 17


def test_allocation_review_is_one_editable_row_per_fixture():
    engine = db()
    open_and_vote_all(engine, [1, 2])
    propose(engine)
    review = allocation_review(engine)
    assert len(review) == 2
    assert list({key: None for key in review[0] if key != "fixture_id"}) == REVIEW_COLUMNS
    assert all(row["Umpire 1"] and row["Umpire 2"] and row["Scorer"] for row in review)
    assert all(row["Status"] == "Proposed" for row in review)


def test_allocation_review_row_save_then_confirm_is_audited():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    review = allocation_review(engine)[0]
    original = review["Umpire 1"]
    excluded = {review["Umpire 1"], review["Umpire 2"], review["Scorer"]}
    review["Umpire 1"] = next(
        person["name"] for person in rows(engine, "SELECT name FROM people WHERE active=true AND can_umpire=true ORDER BY name")
        if person["name"] not in excluded
    )
    save_allocation_review(engine, [review], confirm=False)
    saved = rows(engine, """
        SELECT p.name,a.confirmed FROM assignments a JOIN people p ON p.id=a.person_id
        WHERE a.fixture_id=1 AND a.role='umpire_1'
    """)[0]
    assert saved == {"name": review["Umpire 1"], "confirmed": 0}
    assert rows(engine, "SELECT event_type FROM assignment_events WHERE fixture_id=1") == [{"event_type": "MANUAL_CHANGE"}]
    assert confirmed_allocation_board(engine) == []

    save_allocation_review(engine, [review], confirm=True)
    assert rows(engine, "SELECT COUNT(*) n FROM assignments WHERE fixture_id=1 AND confirmed=true AND status='ASSIGNED'")[0]["n"] == 3
    assert rows(engine, "SELECT COUNT(*) n FROM assignment_events WHERE fixture_id=1 AND event_type='CONFIRMED'")[0]["n"] == 3
    assert confirmed_allocation_board(engine)[0]["Umpire 1"] == review["Umpire 1"]
    manual_event = rows(engine, """
        SELECT old.name person,new.name replacement,e.event_type
        FROM assignment_events e JOIN people old ON old.id=e.person_id
        JOIN people new ON new.id=e.replacement_person_id
        WHERE e.event_type='MANUAL_CHANGE'
    """)[0]
    assert manual_event == {"person": original, "replacement": review["Umpire 1"], "event_type": "MANUAL_CHANGE"}


def test_confirm_all_proposed_and_final_exports_are_confirmed_only():
    engine = db()
    open_and_vote_all(engine, [1, 2])
    propose(engine)
    assert confirmed_allocation_board(engine) == []
    assert allocation_message(engine).strip().endswith("Final Allocation*")
    assert confirm_all_proposed(engine) == 2
    final = confirmed_allocation_board(engine)
    assert len(final) == 2
    csv_text = allocation_board_csv(final)
    assert csv_text.splitlines()[0] == ",".join(BOARD_COLUMNS)
    assert final[0]["Umpire 1"] in csv_text
    message = allocation_message(engine)
    assert "SGIA Umpires" in message
    assert "Changi Risers vs Black Panthers" in message
    assert "Umpire 1:" in message and "Umpire 2:" in message and "Scorer:" in message


def test_output_excludes_incomplete_or_replacement_required_assignments():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    with engine.begin() as connection:
        connection.execute(text("UPDATE assignments SET confirmed=true WHERE fixture_id=1"))
        connection.execute(text("UPDATE assignments SET status='REPLACEMENT_REQUIRED' WHERE fixture_id=1 AND role='umpire_1'"))
    assert confirmed_allocation_board(engine) == []
    assert "Changi Risers" not in allocation_message(engine)


def test_admin_allocation_review_has_desktop_and_mobile_contracts():
    source = Path("app.py").read_text(encoding="utf-8")
    assert 'key="desktop_allocation_review"' in source
    assert 'key="mobile_allocation_review"' in source
    assert "@media(max-width:700px)" in source
    assert ".st-key-desktop_allocation_review{display:none!important}" in source
    assert "Save row" in source and "Confirm row" in source
    assert "Generate proposed allocation" in source and "Regenerate unconfirmed only" in source
    assert "Final Confirmed Allocation" in source
    assert "Download PNG" in source and "Download CSV" in source and "Copy WhatsApp Message" in source


def test_public_route_has_no_admin_navigation():
    assert not is_admin_request({})
    assert is_admin_request({"admin": "1"})
    source = Path("app.py").read_text(encoding="utf-8")
    public_body = source.split("def render_public", 1)[1].split("def admin_authenticated", 1)[0]
    assert "Admin" not in public_body
    assert "ADMIN_MODE = is_admin_request(st.query_params)" in source
    assert "if ADMIN_MODE:" in source


def test_public_page_runtime_has_no_admin_tabs(monkeypatch, tmp_path):
    monkeypatch.setenv("SGIA_SQLITE_PATH", str(tmp_path / "public.db"))
    page = AppTest.from_file(Path(__file__).parents[1] / "app.py").run(timeout=20)
    assert not list(page.exception)
    assert len(page.tabs) == 0
    assert not any("Admin" in getattr(item, "value", "") for item in page.markdown)
    rendered = " ".join(getattr(item, "value", "") for item in page.markdown)
    assert "Provide your availability for this weekend" in rendered
    assert "Select your name to view the available slots." in rendered


def test_public_brand_asset_and_mobile_layout_are_required():
    root = Path(__file__).parents[1]
    logo = root / "assets" / "sgia-logo.png"
    assert logo.is_file()
    assert logo.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
    source = (root / "app.py").read_text(encoding="utf-8")
    assert "max-width:500px!important" in source
    assert "brand-badge" not in source
    assert 'raise FileNotFoundError("Required SGIA logo is missing' in source


def test_admin_query_route_still_renders_login(monkeypatch, tmp_path):
    monkeypatch.setenv("SGIA_SQLITE_PATH", str(tmp_path / "admin.db"))
    monkeypatch.setenv("SGIA_ADMIN_PIN", "test-only-pin")
    page = AppTest.from_file(Path(__file__).parents[1] / "app.py")
    page.query_params["admin"] = "1"
    page.run(timeout=20)
    assert not list(page.exception)
    assert any(item.label == "Admin PIN" for item in page.text_input)
    next(item for item in page.text_input if item.label == "Admin PIN").set_value("test-only-pin")
    next(item for item in page.button if item.label == "Sign in").click()
    page.run(timeout=20)
    assert not list(page.exception)
    navigation = next(item for item in page.radio if item.label == "Admin section")
    assert set(navigation.options) == {
        "Quick Admin", "Open Poll", "Fixtures", "Allocations", "Allocation Board",
        "People", "Season Workload", "History", "Output",
    }
    navigation.set_value("Allocations")
    page.run(timeout=20)
    assert not list(page.exception)
    assert any(item.value == "Allocation Review" for item in page.subheader)
    next(item for item in page.radio if item.label == "Admin section").set_value("Fixtures")
    page.run(timeout=20)
    next(item for item in page.radio if item.label == "Fixture action").set_value("Bulk Import")
    page.run(timeout=20)
    assert not list(page.exception)
    assert any(item.label == "UPLOAD FIXTURE FILE" for item in page.file_uploader)
    download_labels = {item.label for item in page.get("download_button")}
    assert {"DOWNLOAD CSV TEMPLATE", "DOWNLOAD EXCEL TEMPLATE"} <= download_labels


def test_admin_navigation_renders_only_the_selected_section():
    source = (Path(__file__).parents[1] / "app.py").read_text(encoding="utf-8")
    assert 'admin_section = st.radio(' in source
    assert 'st.tabs(["Quick Admin", "Open Poll"' not in source
    for section in ("Quick Admin", "Open Poll", "Fixtures", "Allocations", "Allocation Board", "People", "Season Workload", "History", "Output"):
        assert f'if admin_section == "{section}":' in source


def test_confirmed_allocation_message():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    with engine.begin() as connection:
        connection.execute(text("UPDATE assignments SET confirmed=true"))
    message = allocation_message(engine)
    assert "SGIA Umpires" in message
    assert "Umpire 1:" in message
    assert "Umpire 2:" in message


def test_published_duty_withdrawal_requires_replacement():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    assignment = rows(engine, "SELECT fixture_id,role,person_id FROM assignments WHERE fixture_id=1 ORDER BY role LIMIT 1")[0]
    with engine.begin() as connection:
        connection.execute(text("UPDATE assignments SET confirmed=true WHERE fixture_id=1"))
        connection.execute(text("UPDATE fixtures SET availability_open=false,poll_state='PUBLISHED' WHERE id=1"))
    assert withdraw_assignment(engine, assignment["fixture_id"], assignment["role"], assignment["person_id"])
    assert rows(engine, "SELECT status FROM assignments WHERE fixture_id=:fixture_id AND role=:role", assignment) == [{"status": "REPLACEMENT_REQUIRED"}]
    assert rows(engine, "SELECT event_type,person_id FROM assignment_events") == [{"event_type": "WITHDRAWN", "person_id": assignment["person_id"]}]
    assert not withdraw_assignment(engine, assignment["fixture_id"], assignment["role"], assignment["person_id"])
    replacement = suggest_replacement(engine, assignment["fixture_id"], assignment["role"])
    assert replacement and replacement["id"] != assignment["person_id"]
    assert rows(engine, "SELECT person_id,confirmed,status FROM assignments WHERE fixture_id=:fixture_id AND role=:role", assignment) == [
        {"person_id": replacement["id"], "confirmed": 0, "status": "ASSIGNED"}
    ]
    events = rows(engine, "SELECT event_type,person_id,replacement_person_id FROM assignment_events ORDER BY id")
    assert events[0]["event_type"] == "WITHDRAWN"
    assert events[1] == {"event_type": "REPLACED", "person_id": assignment["person_id"], "replacement_person_id": replacement["id"]}


def test_manual_replacement_records_replaced_event():
    engine = db()
    open_and_vote_all(engine, [1])
    propose(engine)
    assignment = rows(engine, "SELECT fixture_id,role,person_id FROM assignments WHERE fixture_id=1 ORDER BY role LIMIT 1")[0]
    replacement = rows(engine, "SELECT id FROM people WHERE active=true AND id<>:person ORDER BY id LIMIT 1", {"person": assignment["person_id"]})[0]["id"]
    replace_assignment(engine, assignment["fixture_id"], assignment["role"], replacement)
    assert rows(engine, "SELECT event_type,person_id,replacement_person_id FROM assignment_events") == [
        {"event_type": "REPLACED", "person_id": assignment["person_id"], "replacement_person_id": replacement}
    ]


def test_fixture_add_edit_and_duplicate_safe_bulk_import():
    engine = db()
    assert add_fixture(engine, "05-Sep-2026", "11:00 AM", "Team A", "Team B")
    assert not add_fixture(engine, "05-Sep-2026", "11:00 AM", "Duplicate", "Ignored")
    fixture = rows(engine, "SELECT id FROM fixtures WHERE home_team='Team A'")[0]
    assert edit_fixture(engine, fixture["id"], "05-Sep-2026", "12:00 PM", "Team Alpha", "Team Beta")
    edited = rows(engine, "SELECT home_team,away_team FROM fixtures WHERE id=:id", fixture)[0]
    assert edited == {"home_team": "Team Alpha", "away_team": "Team Beta"}
    pasted = """Day | Date | Time | TEAM 1 | TEAM 2
Saturday | 05-Sep-2026 | 12:00 PM | Duplicate | Existing
Sunday | 06-Sep-2026 | 3:00 PM | Team C | Team D
Sunday | 06-Sep-2026 | 3:00 PM | Duplicate In Paste | Team E"""
    preview = preview_bulk_fixtures(engine, pasted)
    assert [item["Status"] for item in preview] == ["DUPLICATE", "READY", "DUPLICATE"]
    assert import_bulk_fixtures(engine, preview) == 1
    assert import_bulk_fixtures(engine, preview) == 0


def test_fixture_csv_and_excel_templates_are_real_and_formatted():
    csv_text = fixture_template_csv().decode("utf-8-sig")
    assert csv_text.splitlines()[0] == ",".join(FIXTURE_COLUMNS)
    assert "Changi Risers" in csv_text and "Knights United" in csv_text

    workbook = load_workbook(BytesIO(fixture_template_xlsx()))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == FIXTURE_COLUMNS
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.column_dimensions["D"].width >= 20
    assert isinstance(sheet["B2"].value, (date, datetime))
    assert isinstance(sheet["C2"].value, (time, datetime))
    assert sheet["A2"].comment and "Example row" in sheet["A2"].comment.text


def test_csv_upload_accepts_alias_headers_optional_day_and_24_hour_time():
    engine = db()
    uploaded = (
        "Date,Time,Home Team,Away Team\n"
        "15/09/2026,19:00,Warriors,Legends\n"
        "2026-09-16,7:30 AM,Team Alpha,Team Beta\n"
    ).encode("utf-8")
    preview = preview_fixture_file(engine, "schedule.csv", uploaded)
    assert [item["Status"] for item in preview] == ["READY", "READY"]
    assert preview[0]["Day"] == "Tuesday"
    assert preview[0]["starts_at"] == datetime(2026, 9, 15, 19, 0)
    assert preview[1]["starts_at"] == datetime(2026, 9, 16, 7, 30)


def test_excel_style_two_digit_year_text_is_accepted_as_current_century():
    engine = db()
    pasted = """Day\tDate\tTime\tTEAM 1\tTEAM 2
Tuesday\t01-Sep-26\t7:00 PM\tTamil Titans\tCenturions
Wednesday\t02-Sep-26\t7:00 PM\tKarunadu\tCW Storm
Thursday\t03/09/26\t7:00 PM\tKnights United\tRoyal Star"""
    preview = preview_bulk_fixtures(engine, pasted)
    assert [item["Status"] for item in preview] == ["READY", "READY", "READY"]
    assert [item["starts_at"] for item in preview] == [
        datetime(2026, 9, 1, 19, 0),
        datetime(2026, 9, 2, 19, 0),
        datetime(2026, 9, 3, 19, 0),
    ]
    assert [item["Date"] for item in preview] == ["01-Sep-2026", "02-Sep-2026", "03-Sep-2026"]


def test_xlsx_upload_accepts_excel_date_and_time_cells_without_timezone_shift():
    engine = db()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Date", "Time", "Team 1", "Team 2"])
    sheet.append([date(2026, 9, 17), time(15, 0), "Excel XI", "Spreadsheet XI"])
    payload = BytesIO()
    workbook.save(payload)
    preview = preview_fixture_file(engine, "schedule.xlsx", payload.getvalue())
    assert len(preview) == 1 and preview[0]["Status"] == "READY"
    assert preview[0]["Day"] == "Thursday"
    assert preview[0]["Date"] == "17-Sep-2026"
    assert preview[0]["Time"] == "3:00 PM"
    assert preview[0]["starts_at"] == datetime(2026, 9, 17, 15, 0)
    assert preview[0]["starts_at"].tzinfo is None


def test_pipe_tab_and_csv_paste_share_preview_and_do_not_split_team_spaces():
    engine = db()
    pipe_preview = preview_bulk_fixtures(
        engine, "Day | Date | Time | TEAM 1 | TEAM 2\nFriday | 18-Sep-2026 | 7:00 PM | United Warriors | Royal Legends"
    )
    tab_preview = preview_bulk_fixtures(
        engine, "Tuesday\t22-Sep-2026\t7:00 PM\tWarriors United\tLegends United"
    )
    csv_preview = preview_bulk_fixtures(
        engine, "Date,Time,Team 1,Team 2\n2026-09-23,19:00,Comma Warriors,Comma Legends"
    )
    assert pipe_preview[0]["TEAM 1"] == "United Warriors"
    assert tab_preview[0]["Status"] == "READY"
    assert tab_preview[0]["TEAM 2"] == "Legends United"
    assert csv_preview[0]["Status"] == "READY"
    unsafe_spaces = preview_bulk_fixtures(engine, "23-Sep-2026 7:00 PM Team With Spaces Other Team")
    assert unsafe_spaces[0]["Status"] == "ERROR"
    assert "spaces cannot safely separate" in unsafe_spaces[0]["Error"]


def test_mixed_batch_classifies_ready_duplicate_error_and_never_overwrites():
    engine = db()
    original = rows(engine, "SELECT id,home_team,away_team FROM fixtures WHERE id=1")[0]
    batch = """Date,Time,Home Team,Away Team
15-Aug-2026,11:00 AM,Overwrite Attempt,Must Not Replace
24-Sep-2026,7:00 PM,Ready Team,Ready Opponent
24-Sep-2026,7:00 PM,Repeated Ready Time,Also Duplicate
not-a-date,3:00 PM,Broken,Row"""
    preview = preview_fixture_file(engine, "mixed.csv", batch.encode())
    assert [item["Status"] for item in preview] == ["DUPLICATE", "READY", "DUPLICATE", "ERROR"]
    assert import_bulk_fixtures(engine, preview) == 1
    assert import_bulk_fixtures(engine, preview) == 0
    unchanged = rows(engine, "SELECT id,home_team,away_team FROM fixtures WHERE id=1")[0]
    assert unchanged == original
    assert rows(engine, "SELECT COUNT(*) n FROM fixtures WHERE starts_at='2026-09-24 19:00:00'")[0]["n"] == 1


def test_invalid_file_headers_are_reported_and_current_fixture_exports_are_safe():
    engine = db()
    invalid = preview_fixture_file(engine, "bad.csv", b"Date,Time,Home Team\n2026-09-25,19:00,Only One Team")
    assert invalid[0]["Status"] == "ERROR"
    assert "Missing required column" in invalid[0]["Error"]

    csv_header = current_fixtures_csv(engine).decode("utf-8-sig").splitlines()[0]
    assert csv_header == ",".join(CURRENT_FIXTURE_COLUMNS)
    workbook = load_workbook(BytesIO(current_fixtures_xlsx(engine)), read_only=True)
    assert [cell.value for cell in workbook.active[1]] == CURRENT_FIXTURE_COLUMNS
    assert workbook.active.max_row == 30


def test_fixture_bulk_import_admin_contract_is_visible():
    source = (Path(__file__).parents[1] / "app.py").read_text(encoding="utf-8")
    for label in (
        "Fixture Bulk Import", "DOWNLOAD CSV TEMPLATE", "DOWNLOAD EXCEL TEMPLATE",
        "UPLOAD FIXTURE FILE", "OR Paste from Excel / Text", "IMPORT READY FIXTURES",
        "DOWNLOAD CURRENT FIXTURES CSV", "DOWNLOAD CURRENT FIXTURES EXCEL",
    ):
        assert label in source


def test_existing_schema_migration_preserves_data():
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE people (id INTEGER PRIMARY KEY, name VARCHAR(80) UNIQUE NOT NULL, can_umpire BOOLEAN NOT NULL, can_score BOOLEAN NOT NULL, preferred_role VARCHAR(10) NOT NULL, active BOOLEAN NOT NULL DEFAULT TRUE)"))
        connection.execute(text("CREATE TABLE fixtures (id INTEGER PRIMARY KEY, starts_at TIMESTAMP UNIQUE NOT NULL, home_team VARCHAR(100) NOT NULL, away_team VARCHAR(100) NOT NULL, availability_open BOOLEAN NOT NULL DEFAULT FALSE)"))
        connection.execute(text("CREATE TABLE availability (person_id INTEGER NOT NULL, fixture_id INTEGER NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(person_id,fixture_id))"))
        connection.execute(text("CREATE TABLE assignments (fixture_id INTEGER NOT NULL, role VARCHAR(10) NOT NULL, person_id INTEGER NOT NULL, confirmed BOOLEAN NOT NULL DEFAULT FALSE, reason TEXT, PRIMARY KEY(fixture_id,role))"))
        connection.execute(text("INSERT INTO people VALUES(100,'Legacy Person',true,true,'Either',true)"))
        connection.execute(text("INSERT INTO fixtures VALUES(100,'2026-09-20 10:00:00','Legacy A','Legacy B',true)"))
        connection.execute(text("INSERT INTO availability(person_id,fixture_id) VALUES(100,100)"))
        connection.execute(text("INSERT INTO assignments VALUES(100,'umpire',100,true,'Legacy assignment')"))
    initialize(engine)
    assert rows(engine, "SELECT name,pin_hash FROM people WHERE id=100") == [{"name": "Legacy Person", "pin_hash": None}]
    assert rows(engine, "SELECT poll_state FROM fixtures WHERE id=100") == [{"poll_state": "OPEN"}]
    assert rows(engine, "SELECT person_id,fixture_id FROM availability WHERE person_id=100") == [{"person_id": 100, "fixture_id": 100}]
    assert rows(engine, "SELECT role,status,person_id FROM assignments WHERE fixture_id=100") == [{"role": "umpire_1", "status": "ASSIGNED", "person_id": 100}]
    assert rows(engine, "SELECT match_status FROM fixtures WHERE id=100") == [{"match_status": "SCHEDULED"}]
    assert rows(engine, "SELECT COUNT(*) n FROM fixture_events")[0]["n"] == 0
    assert rows(engine, "SELECT COUNT(*) n FROM poll_submissions WHERE person_id=100")[0]["n"] == 1


def test_match_status_migration_is_idempotent_and_defaults_scheduled():
    engine = db()
    initialize(engine)
    assert rows(engine, "SELECT DISTINCT match_status FROM fixtures") == [{"match_status": "SCHEDULED"}]
    assert rows(engine, "SELECT COUNT(*) n FROM fixtures")[0]["n"] == 29


def test_completed_workload_counts_distinct_fixture_roles_and_scorer():
    engine = db()
    person = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=1"))
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(1,'umpire_1',:person,true,'ASSIGNED')"), {"person": person})
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(1,'scorer',:person,true,'ASSIGNED')"), {"person": person})
    set_match_status(engine, 1, "COMPLETED")
    stats = season_workload(engine, person)[0]
    assert stats["games_completed"] == 1
    assert stats["umpire_completed"] == 1
    assert stats["scorer_completed"] == 1
    assert stats["total_assigned"] == 1


def test_cancelled_and_directly_restored_completed_fixture_reverse_workload():
    engine = db()
    person = rows(engine, "SELECT id FROM people WHERE name='Malo'")[0]["id"]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=1"))
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(1,'umpire_2',:person,true,'ASSIGNED')"), {"person": person})
    set_match_status(engine, 1, "COMPLETED")
    assert season_workload(engine, person)[0]["games_completed"] == 1
    assert season_workload(engine, person)[0]["umpire_completed"] == 1
    set_match_status(engine, 1, "SCHEDULED")
    assert season_workload(engine, person)[0]["games_completed"] == 0
    set_match_status(engine, 1, "COMPLETED")
    set_match_status(engine, 1, "CANCELLED")
    assert season_workload(engine, person)[0]["games_completed"] == 0
    assert [event["action"] for event in fixture_history(engine)] == ["MATCH_CANCELLED", "MATCH_COMPLETED", "MATCH_RESTORED", "MATCH_COMPLETED"]


def test_unpublished_fixture_cannot_be_completed():
    engine = db()
    try:
        set_match_status(engine, 1, "COMPLETED")
        assert False, "Expected unpublished completion to be rejected"
    except ValueError as error:
        assert "Publish" in str(error)
    assert rows(engine, "SELECT match_status FROM fixtures WHERE id=1") == [{"match_status": "SCHEDULED"}]


def test_upcoming_requires_future_published_confirmed_active_assignment():
    engine = db()
    people = rows(engine, "SELECT id FROM people ORDER BY id LIMIT 3")
    assert add_fixture(engine, "01-Jan-2099", "11:00 AM", "Future Published", "Future Opponent")
    assert add_fixture(engine, "02-Jan-2099", "11:00 AM", "Future Unpublished", "Future Opponent")
    assert add_fixture(engine, "03-Jan-2099", "11:00 AM", "Future Unconfirmed", "Future Opponent")
    future_ids = [item["id"] for item in rows(engine, "SELECT id FROM fixtures WHERE starts_at>='2099-01-01' ORDER BY starts_at")]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=:fixture"), {"fixture": future_ids[0]})
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(:fixture,'umpire_1',:person,true,'ASSIGNED')"), {"fixture": future_ids[0], "person": people[0]["id"]})
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(:fixture,'umpire_1',:person,true,'ASSIGNED')"), {"fixture": future_ids[1], "person": people[1]["id"]})
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(:fixture,'umpire_1',:person,false,'ASSIGNED')"), {"fixture": future_ids[2], "person": people[2]["id"]})
    assert season_workload(engine, people[0]["id"])[0]["upcoming_duties"] == 1
    assert season_workload(engine, people[1]["id"])[0]["upcoming_duties"] == 0
    assert season_workload(engine, people[2]["id"])[0]["upcoming_duties"] == 0


def test_withdrawal_does_not_count_and_replacement_credits_current_person():
    engine = db()
    original, replacement = rows(engine, "SELECT id FROM people ORDER BY id LIMIT 2")
    assert add_fixture(engine, "04-Jan-2099", "11:00 AM", "Future Withdrawal", "Future Opponent")
    fixture_id = rows(engine, "SELECT id FROM fixtures WHERE starts_at='2099-01-04 11:00:00'")[0]["id"]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=:fixture"), {"fixture": fixture_id})
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(:fixture,'scorer',:person,true,'ASSIGNED')"), {"fixture": fixture_id, "person": original["id"]})
    assert withdraw_assignment(engine, fixture_id, "scorer", original["id"])
    assert season_workload(engine, original["id"])[0]["upcoming_duties"] == 0
    replace_assignment(engine, fixture_id, "scorer", replacement["id"])
    assert season_workload(engine, replacement["id"])[0]["upcoming_duties"] == 1
    assert season_workload(engine, original["id"])[0]["upcoming_duties"] == 0


def test_workload_orders_least_completed_then_total_then_name():
    engine = db()
    people = season_workload(engine)
    assert [person["name"] for person in people] == sorted(person["name"] for person in people)
    person = people[-1]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=1"))
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(1,'umpire_1',:person,true,'ASSIGNED')"), {"person": person["person_id"]})
    set_match_status(engine, 1, "COMPLETED")
    assert season_workload(engine)[-1]["person_id"] == person["person_id"]


def test_allocator_uses_lower_season_workload_only_as_tiebreaker():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO people(name,can_umpire,can_score,preferred_role,active) VALUES('Fair High',true,true,'Either',true),('Fair Low',true,true,'Either',true)"))
        high = connection.execute(text("SELECT id FROM people WHERE name='Fair High'")).scalar_one()
        low = connection.execute(text("SELECT id FROM people WHERE name='Fair Low'")).scalar_one()
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED',match_status='COMPLETED' WHERE id=2"))
        connection.execute(text("INSERT INTO assignments(fixture_id,role,person_id,confirmed,status) VALUES(2,'umpire_1',:high,true,'ASSIGNED')"), {"high": high})
        connection.execute(text("INSERT INTO availability(person_id,fixture_id) VALUES(:high,12),(:low,12)"), {"high": high, "low": low})
    propose(engine)
    chosen = rows(engine, "SELECT person_id,reason FROM assignments WHERE fixture_id=12 AND role='umpire_1'")[0]
    assert chosen["person_id"] == low
    assert "season workload" in chosen["reason"].lower()


def test_live_poll_counts_active_zero_selection_response_and_nonresponders():
    engine = db()
    person = rows(engine, "SELECT id,name FROM people WHERE active=true ORDER BY name LIMIT 1")[0]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='OPEN',availability_open=true WHERE id IN (1,2)"))
    save_vote(engine, person["id"], [])
    monitor = live_poll_monitor(engine)
    assert monitor["active_volunteers"] == 16
    assert monitor["responded"] == 1
    assert monitor["open_matches"] == 2
    assert person["name"] not in monitor["yet_to_respond"]
    assert len(monitor["yet_to_respond"]) == 15
    assert rows(engine, "SELECT COUNT(*) n FROM availability WHERE person_id=:person", {"person": person["id"]})[0]["n"] == 0


def test_live_poll_updated_submission_remains_one_respondent():
    engine = db()
    person = rows(engine, "SELECT id FROM people WHERE active=true ORDER BY id LIMIT 1")[0]
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='OPEN',availability_open=true WHERE id IN (1,2)"))
    save_vote(engine, person["id"], [1])
    save_vote(engine, person["id"], [2])
    assert live_poll_monitor(engine)["responded"] == 1
    assert rows(engine, "SELECT COUNT(*) n FROM poll_submissions WHERE person_id=:id", person)[0]["n"] == 1


def test_live_poll_day_and_match_unique_availability_capabilities():
    engine = db()
    people = rows(engine, "SELECT id,name,can_umpire,can_score FROM people WHERE active=true ORDER BY id LIMIT 3")
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='OPEN',availability_open=true WHERE id IN (1,2)"))
    save_vote(engine, people[0]["id"], [1, 2])
    save_vote(engine, people[1]["id"], [1])
    save_vote(engine, people[2]["id"], [1])
    monitor = live_poll_monitor(engine)
    assert len(monitor["days"]) == 1
    day = monitor["days"][0]
    assert day["label"] == "Saturday 15 Aug"
    assert day["unique_available"] == 3
    assert len(day["matches"]) == 2
    first, second = day["matches"]
    assert first["available_count"] == 3
    assert second["available_count"] == 1
    assert first["umpire_count"] == len({person["id"] for person in people if person["can_umpire"]})
    assert first["scorer_count"] == len({person["id"] for person in people if person["can_score"]})


def test_live_poll_coverage_requires_three_distinct_role_eligible_people():
    covered = [
        {"id": 1, "can_umpire": True, "can_score": True},
        {"id": 2, "can_umpire": True, "can_score": False},
        {"id": 3, "can_umpire": False, "can_score": True},
    ]
    scorer_shortage = [{"id": value, "can_umpire": True, "can_score": False} for value in range(1, 4)]
    umpire_shortage = [
        {"id": 1, "can_umpire": True, "can_score": True},
        {"id": 2, "can_umpire": False, "can_score": True},
        {"id": 3, "can_umpire": False, "can_score": True},
    ]
    assert coverage_status(covered) == "COVERED"
    assert coverage_status(scorer_shortage) == "SCORER COVERAGE NEEDED"
    assert coverage_status(umpire_shortage) == "UMPIRE COVERAGE NEEDED"
    assert coverage_status(covered[:2]) == "NEED MORE"


def test_live_poll_excludes_non_open_fixtures_and_uses_singapore_wall_clock():
    engine = db()
    with engine.begin() as connection:
        connection.execute(text("UPDATE fixtures SET poll_state='OPEN',availability_open=true WHERE id=1"))
        connection.execute(text("UPDATE fixtures SET poll_state='FROZEN' WHERE id=2"))
        connection.execute(text("UPDATE fixtures SET poll_state='CLOSED' WHERE id=3"))
        connection.execute(text("UPDATE fixtures SET poll_state='PUBLISHED' WHERE id=4"))
    monitor = live_poll_monitor(engine)
    assert monitor["open_matches"] == 1
    assert monitor["days"][0]["label"] == "Saturday 15 Aug"
    assert monitor["days"][0]["matches"][0]["time"] == "11:00 AM"


def test_live_poll_mobile_admin_contract_is_present():
    source = Path("app.py").read_text(encoding="utf-8")
    assert "LIVE POLL" in source
    assert "REFRESH LIVE POLL" in source
    assert "Yet to Respond" in source
    assert "@media(max-width:700px)" in source
    assert "live_poll_monitor(db)" in source


def test_poll_cycle_is_persistent_and_revisions_track_fixture_changes():
    engine = db()
    first_cycle = update_open_slots(engine, [1, 2])
    assert first_cycle
    assert rows(engine, "SELECT revision,status FROM poll_cycles WHERE id=:id", {"id": first_cycle}) == [{"revision": 1, "status": "OPEN"}]
    assert rows(engine, "SELECT DISTINCT poll_cycle_id FROM fixtures WHERE id IN (1,2)") == [{"poll_cycle_id": first_cycle}]

    set_fixture_poll_state(engine, 3, "OPEN")
    assert rows(engine, "SELECT poll_cycle_id FROM fixtures WHERE id=3")[0]["poll_cycle_id"] == first_cycle
    assert rows(engine, "SELECT revision FROM poll_cycles WHERE id=:id", {"id": first_cycle})[0]["revision"] == 2
    set_fixture_poll_state(engine, 2, "CLOSED")
    assert rows(engine, "SELECT revision FROM poll_cycles WHERE id=:id", {"id": first_cycle})[0]["revision"] == 3

    transition_poll_state(engine, "OPEN", "FROZEN")
    assert rows(engine, "SELECT status,frozen_at IS NOT NULL frozen FROM poll_cycles WHERE id=:id", {"id": first_cycle}) == [{"status": "FROZEN", "frozen": 1}]
    transition_poll_state(engine, "FROZEN", "PUBLISHED")
    second_cycle = update_open_slots(engine, [12, 13])
    assert second_cycle != first_cycle


def test_working_allocation_and_every_export_are_scoped_to_current_cycle():
    engine = db()
    people = rows(engine, "SELECT id FROM people WHERE active=true ORDER BY id LIMIT 3")

    old_cycle = update_open_slots(engine, [1])
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO availability(person_id,fixture_id) SELECT id,1 FROM people WHERE active=true"))
    transition_poll_state(engine, "OPEN", "FROZEN")
    propose(engine)
    assert confirm_all_proposed(engine) == 1
    old_assignment_count = rows(engine, "SELECT COUNT(*) n FROM assignments WHERE fixture_id=1")[0]["n"]
    transition_poll_state(engine, "FROZEN", "PUBLISHED")

    # A future assigned fixture outside any active cycle must never leak into
    # the current board or exports merely because it is in the future.
    with engine.begin() as connection:
        for role, person in zip(("umpire_1", "umpire_2", "scorer"), people):
            connection.execute(text("""
                INSERT INTO assignments(fixture_id,role,person_id,confirmed,reason,status)
                VALUES(20,:role,:person,true,'Outside cycle','ASSIGNED')
            """), {"role": role, "person": person["id"]})

    assert allocation_board(engine) == []
    assert confirmed_allocation_board(engine) == []
    assert rows(engine, "SELECT COUNT(*) n FROM assignments WHERE fixture_id=1")[0]["n"] == old_assignment_count

    current_cycle = update_open_slots(engine, [12])
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO availability(person_id,fixture_id) SELECT id,12 FROM people WHERE active=true"))
    transition_poll_state(engine, "OPEN", "FROZEN")
    propose(engine)
    assert confirm_all_proposed(engine) == 1

    summary = current_allocation_summary(engine)
    assert summary["id"] == current_cycle
    assert summary["status"] == "FROZEN"
    assert summary["match_count"] == 1
    current_rows = confirmed_allocation_board(engine)
    assert [item["fixture_id"] for item in current_rows] == [12]
    current_fixture = rows(engine, "SELECT home_team,away_team FROM fixtures WHERE id=12")[0]
    old_fixture = rows(engine, "SELECT home_team FROM fixtures WHERE id=1")[0]
    outside_fixture = rows(engine, "SELECT home_team FROM fixtures WHERE id=20")[0]

    csv_output = allocation_board_csv(current_rows)
    message = allocation_message(engine)
    assert current_fixture["home_team"] in csv_output and current_fixture["home_team"] in message
    assert old_fixture["home_team"] not in csv_output and old_fixture["home_team"] not in message
    assert outside_fixture["home_team"] not in csv_output and outside_fixture["home_team"] not in message
    png = Image.open(BytesIO(allocation_board_png(current_rows)))
    assert png.height == 116 + 72 + 86 + 34  # exactly one exported match row

    published_ids = {cycle["id"] for cycle in published_allocation_cycles(engine)}
    assert old_cycle in published_ids
    assert rows(engine, "SELECT COUNT(*) n FROM assignments WHERE fixture_id=1")[0]["n"] == old_assignment_count

    transition_poll_state(engine, "FROZEN", "PUBLISHED")
    assert allocation_board(engine) == []
    next_cycle = update_open_slots(engine, [20])
    assert next_cycle not in {old_cycle, current_cycle}
    assert current_allocation_summary(engine)["id"] == next_cycle
    assert [item["fixture_id"] for item in allocation_board(engine)] == [20]


def test_frozen_cycle_takes_priority_then_open_cycle_becomes_current_after_publish():
    engine = db()
    frozen_cycle = update_open_slots(engine, [1])
    transition_poll_state(engine, "OPEN", "FROZEN")
    open_cycle = update_open_slots(engine, [12])
    assert open_cycle != frozen_cycle
    assert current_allocation_summary(engine)["id"] == frozen_cycle
    transition_poll_state(engine, "FROZEN", "PUBLISHED")
    assert current_allocation_summary(engine)["id"] == open_cycle


def test_view_and_submission_evidence_is_durable_and_snapshot_complete():
    engine = db()
    person = rows(engine, "SELECT id FROM people ORDER BY id LIMIT 1")[0]["id"]
    cycle = update_open_slots(engine, [1, 2, 3])
    record_poll_view(engine, person, "OPEN", "Mobile", "Chrome")
    record_poll_view(engine, person, "OPEN", "Mobile", "Chrome")
    assert rows(engine, "SELECT COUNT(*) n FROM poll_activity_events WHERE event_type='VIEWED'")[0]["n"] == 1
    viewed = next(item for item in response_audit(engine, cycle)["people"] if item["person_id"] == person)
    assert viewed["Status"] == "VIEWED — NOT SUBMITTED"

    save_vote(engine, person, [1, 3], "Mobile", "Chrome")
    save_vote(engine, person, [2], "Mobile", "Chrome")
    audited = next(item for item in response_audit(engine, cycle)["people"] if item["person_id"] == person)
    assert audited["Status"] == "UPDATED"
    assert audited["Submission Count"] == 2
    assert audited["Selected Count"] == 1
    assert [item["id"] for item in audited["Current Selections"]] == [2]
    history = activity_history(engine, cycle, person)
    assert [event["event_type"] for event in history] == ["VIEWED", "SUBMITTED", "UPDATED"]
    assert [item["id"] for item in history[1]["selections"]] == [1, 3]
    assert [item["id"] for item in history[2]["selections"]] == [2]


def test_zero_selection_is_response_and_revision_only_adds_needs_review():
    engine = db()
    person = rows(engine, "SELECT id,name FROM people ORDER BY id LIMIT 1")[0]
    cycle = update_open_slots(engine, [1, 2])
    save_vote(engine, person["id"], [])
    initial = live_poll_monitor(engine)
    assert initial["responded"] == 1
    assert initial["responded_zero"] == [person["name"]]
    assert initial["needs_review"] == []

    update_open_slots(engine, [1, 2, 3])
    changed = live_poll_monitor(engine)
    assert changed["responded"] == 1
    assert person["name"] not in changed["yet_to_respond"]
    assert changed["needs_review"] == [person["name"]]
    audited = next(item for item in response_audit(engine, cycle)["people"] if item["person_id"] == person["id"])
    assert audited["Status"] == "RESPONDED — NOT AVAILABLE"
    assert audited["Needs Review"] is True


def test_coarse_client_parser_and_schema_store_no_sensitive_client_data():
    chrome_mobile = "Mozilla/5.0 (Linux; Android 14; Pixel) AppleWebKit/537.36 Chrome/126.0 Mobile Safari/537.36"
    assert coarse_client(chrome_mobile) == ("Mobile", "Chrome")
    assert coarse_client("Mozilla/5.0 (iPad; CPU OS 17_5) Version/17.5 Mobile Safari/604.1") == ("Tablet", "Safari")
    assert coarse_client("Mozilla/5.0 (Windows NT 10.0) Edg/126.0") == ("Desktop", "Edge")
    assert coarse_client("Mozilla/5.0 (X11; Linux x86_64) Firefox/127.0") == ("Desktop", "Firefox")
    engine = db()
    columns = {item["name"].lower() for table in ("poll_participation", "poll_activity_events") for item in __import__("sqlalchemy").inspect(engine).get_columns(table)}
    assert not {"ip", "ip_address", "location", "latitude", "longitude", "user_agent", "fingerprint"} & columns


def test_response_audit_csv_and_admin_only_contract():
    engine = db()
    person = rows(engine, "SELECT id FROM people ORDER BY id LIMIT 1")[0]["id"]
    cycle = update_open_slots(engine, [1])
    save_vote(engine, person, [1], "Desktop", "Firefox")
    audit = response_audit(engine, cycle)
    exported = response_audit_csv(audit)
    assert exported.startswith("Person,Status,First Viewed SGT,First Submitted SGT")
    assert "Desktop,Firefox,LIVE" in exported
    source = Path("app.py").read_text(encoding="utf-8")
    public_body = source.split("def render_public", 1)[1].split("def admin_authenticated", 1)[0]
    admin_body = source.split("def render_admin", 1)[1]
    assert "RESPONSE AUDIT" not in public_body
    assert "DOWNLOAD RESPONSE AUDIT CSV" in admin_body


def test_legacy_evidence_never_fabricates_view_history():
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE people (id INTEGER PRIMARY KEY, name VARCHAR(80) UNIQUE NOT NULL, can_umpire BOOLEAN NOT NULL, can_score BOOLEAN NOT NULL, preferred_role VARCHAR(10) NOT NULL, active BOOLEAN NOT NULL DEFAULT TRUE)"))
        connection.execute(text("CREATE TABLE fixtures (id INTEGER PRIMARY KEY, starts_at TIMESTAMP UNIQUE NOT NULL, home_team VARCHAR(100) NOT NULL, away_team VARCHAR(100) NOT NULL, availability_open BOOLEAN NOT NULL DEFAULT FALSE)"))
        connection.execute(text("CREATE TABLE availability (person_id INTEGER NOT NULL, fixture_id INTEGER NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(person_id,fixture_id))"))
        connection.execute(text("CREATE TABLE assignments (fixture_id INTEGER NOT NULL, role VARCHAR(10) NOT NULL, person_id INTEGER NOT NULL, confirmed BOOLEAN NOT NULL DEFAULT FALSE, reason TEXT, PRIMARY KEY(fixture_id,role))"))
        connection.execute(text("INSERT INTO people VALUES(100,'Legacy Audit',true,true,'Either',true)"))
        connection.execute(text("INSERT INTO fixtures VALUES(100,'2026-09-20 10:00:00','Legacy A','Legacy B',true)"))
        connection.execute(text("INSERT INTO availability(person_id,fixture_id) VALUES(100,100)"))
    initialize(engine)
    imported = rows(engine, "SELECT first_viewed_at,evidence_source FROM poll_participation WHERE person_id=100")[0]
    assert imported == {"first_viewed_at": None, "evidence_source": "LEGACY_BACKFILL"}
    assert rows(engine, "SELECT event_type FROM poll_activity_events WHERE person_id=100") == [{"event_type": "LEGACY_IMPORTED"}]
